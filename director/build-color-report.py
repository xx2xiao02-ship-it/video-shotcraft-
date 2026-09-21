# -*- coding: utf-8 -*-
"""把 color-scan.json 整理成色彩参数分析表 xlsx"""
import json, os, sys

SRC = r"E:\video-shotcraft\director\color-scan.json"
OUT = r"E:\video-shotcraft\director\色彩参数扫描表.xlsx"

with open(SRC, encoding="utf-8") as f:
    rows = json.load(f)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()

# ---------- Sheet1: 汇总统计 ----------
ws = wb.active
ws.title = "汇总"

ws["A1"] = "video-shotcraft 模板色彩参数扫描汇总"
ws["A1"].font = Font(name="微软雅黑", size=14, bold=True, color="FFFFFF")
ws["A1"].fill = PatternFill("solid", fgColor="17365D")
ws.merge_cells("A1:F1")
ws.row_dimensions[1].height = 28

stats = [
    ("总组件数", len(rows)),
    ("引用 G 调色板（共享）", sum(1 for r in rows if r["useG"])),
    ("使用 DesignStage（统一背景容器）", sum(1 for r in rows if r["useDesignStage"])),
    ("纯硬编码（不引用任何共享调色板）", sum(1 for r in rows if not r["useG"] and not r["useDesignStage"])),
    ("硬编码 hex 总数", sum(r["hexCount"] for r in rows)),
    ("rgba 总数", sum(r["rgbaCount"] for r in rows)),
]

headers = ["指标", "数量", "占比", "", "", ""]
for i, h in enumerate(headers, 1):
    c = ws.cell(row=2, column=i, value=h)
    c.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="5B9BD5")
    c.alignment = Alignment(horizontal="center", vertical="center")

for i, (name, val) in enumerate(stats, 3):
    ws.cell(row=i, column=1, value=name).font = Font(name="微软雅黑")
    ws.cell(row=i, column=2, value=val).font = Font(name="微软雅黑")
    pct = f"{val/len(rows)*100:.1f}%" if i-2 <= 3 else ""
    ws.cell(row=i, column=3, value=pct).font = Font(name="微软雅黑")

# ---------- Sheet2: 颜色词频 ----------
ws2 = wb.create_sheet("颜色词频")
ws2["A1"] = "全库颜色使用词频（Top 60）"
ws2["A1"].font = Font(name="微软雅黑", size=14, bold=True, color="FFFFFF")
ws2["A1"].fill = PatternFill("solid", fgColor="17365D")
ws2.merge_cells("A1:C1")

from collections import Counter
counter = Counter()
for r in rows:
    for cat, vals in r["colors"].items():
        for v in vals:
            counter[v] += 1

for i, h in enumerate(["颜色值", "出现次数", "说明"], 1):
    c = ws2.cell(row=2, column=i, value=h)
    c.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="5B9BD5")
    c.alignment = Alignment(horizontal="center", vertical="center")

# 颜色分类说明
color_note = {
    "#0d0f14": "主背景(深)",
    "#1a1d24": "面板背景",
    "#3a3a3a": "灰阶",
    "#5a5a58": "灰阶",
    "#8a8680": "次要文字",
    "#8f8f8d": "灰阶",
    "#d4a853": "金色强调",
    "#4ecdc4": "青色强调",
    "#e8e6e1": "主文字",
    "#ffffff": "白色(纯)",
    "#fff": "白色(简)",
    "#000": "黑色",
    "#111": "深灰",
    "#e6e6e4": "浅灰",
    "#eceef2": "浅灰",
    "#2a2e38": "边框",
    "#c9c9c7": "浅灰",
    "#d8d8d6": "浅灰",
    "#0c0c0c": "深黑",
}
for i, (color, cnt) in enumerate(counter.most_common(60), 3):
    ws2.cell(row=i, column=1, value=color).font = Font(name="Consolas")
    ws2.cell(row=i, column=2, value=cnt).font = Font(name="微软雅黑")
    ws2.cell(row=i, column=3, value=color_note.get(color, "")).font = Font(name="微软雅黑")

# ---------- Sheet3: 组件明细 ----------
ws3 = wb.create_sheet("组件明细")
ws3["A1"] = "组件色彩参数明细（按目录分组）"
ws3["A1"].font = Font(name="微软雅黑", size=14, bold=True, color="FFFFFF")
ws3["A1"].fill = PatternFill("solid", fgColor="17365D")
ws3.merge_cells("A1:L1")

cols = ["组件路径", "分类", "引用G", "DesignStage", "hex数", "rgba数",
        "背景色", "文字色", "边框/阴影", "图形/元素", "强调/渐变", "透明/遮罩"]
for i, h in enumerate(cols, 1):
    c = ws3.cell(row=2, column=i, value=h)
    c.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="5B9BD5")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# 分类：按一级目录
cat_map = {}
for r in rows:
    parts = r["path"].split("/")
    cat = parts[0] if len(parts) > 1 else "other"
    cat_map.setdefault(cat, []).append(r)

row = 3
for cat, items in cat_map.items():
    for r in items:
        c = ws3.cell(row=row, column=1, value=r["path"])
        c.font = Font(name="Consolas", size=9)
        ws3.cell(row=row, column=2, value=cat).font = Font(name="微软雅黑")
        ws3.cell(row=row, column=3, value="✓" if r["useG"] else "").font = Font(name="微软雅黑")
        ws3.cell(row=row, column=4, value="✓" if r["useDesignStage"] else "").font = Font(name="微软雅黑")
        ws3.cell(row=row, column=5, value=r["hexCount"]).font = Font(name="微软雅黑")
        ws3.cell(row=row, column=6, value=r["rgbaCount"]).font = Font(name="微软雅黑")
        colors = r["colors"]
        ws3.cell(row=row, column=7, value=", ".join(colors.get("背景", []))).font = Font(name="Consolas", size=8)
        ws3.cell(row=row, column=8, value=", ".join(colors.get("文字", []))).font = Font(name="Consolas", size=8)
        ws3.cell(row=row, column=9, value=", ".join(colors.get("边框", []) + colors.get("阴影", []))).font = Font(name="Consolas", size=8)
        ws3.cell(row=row, column=10, value=", ".join(colors.get("图形", []) + colors.get("元素", []))).font = Font(name="Consolas", size=8)
        ws3.cell(row=row, column=11, value=", ".join(colors.get("渐变", []))).font = Font(name="Consolas", size=8)
        ws3.cell(row=row, column=12, value=", ".join(colors.get("透明", []))).font = Font(name="Consolas", size=8)
        row += 1

# 列宽
widths1 = {"A": 70, "B": 12, "C": 8, "D": 14, "E": 7, "F": 7, "G": 28, "H": 22, "I": 20, "J": 24, "K": 20, "L": 20}
for col, w in widths1.items():
    ws3.column_dimensions[col].width = w
for col in "ABC":
    ws2.column_dimensions[col].width = 18 if col != "B" else 10
ws.column_dimensions["A"].width = 40
ws.column_dimensions["B"].width = 10
ws.column_dimensions["C"].width = 10

wb.save(OUT)
print(f"✓ 已生成: {OUT}")
print(f"  组件明细行数: {row - 3}")
